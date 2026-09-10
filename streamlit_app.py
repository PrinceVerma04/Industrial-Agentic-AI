"""Streamlit test console for the Sovereign On-Premise Agentic AI Workbench.

Purpose is inspection, not presentation: every tab shows what the system
actually did --- which retrieval mode fired and why, which model was selected
and what it evicted, each tool call with its real arguments and raw output,
and the generated file rendered so you can see its contents rather than trust
a success message. Several of this project's worst defects were invisible in
status lines and obvious the moment the artefact was opened.

    streamlit run streamlit_app.py

This console never opens the graph store or vector index itself - it talks
to the API (app/main.py) over loopback HTTP for everything that touches
them. Kuzu's own Database class does not support a read-only Database
coexisting with a read-write one on the same path (confirmed against the
installed kuzu==0.11.3: "there cannot be multiple Database objects created
with the same database path" when one of them is not read-only) - so the API
process must be the only one that ever opens data/stores/graph. Run the API
first (`make serve` or `make start`); this console is a pure client of it.
"""
from __future__ import annotations

import json
import time
from pathlib import Path

import pandas as pd
import requests
import streamlit as st

from app.core.config import get_settings
from app.graphrag.schema import NODE_TYPES, REL_PAIRS
from app.tools.sandbox import probe, run_python

st.set_page_config(page_title="Sovereign AI Workbench",
                   page_icon="\U0001F512", layout="wide")

S = get_settings()
API = "http://127.0.0.1:8077"


def api_get(path: str, **params):
    r = requests.get(f"{API}{path}", params=params, timeout=180)
    r.raise_for_status()
    return r.json()


def api_post(path: str, **body):
    r = requests.post(f"{API}{path}", json=body, timeout=180)
    r.raise_for_status()
    return r.json()


try:
    health = api_get("/health")
except requests.exceptions.RequestException as e:
    st.error(f"Cannot reach the API at {API}: {type(e).__name__}: {e}")
    st.info("Start it first: `make serve` (or `make start` for both together), "
            "then reload this page.")
    st.stop()


# ------------------------------------------------------------------ sidebar -
with st.sidebar:
    st.title("\U0001F512 Workbench")
    st.caption("SIH26117 \u00b7 MRPL \u00b7 on-premise, air-gapped")

    st.subheader("Sovereignty")
    sov = api_get("/sovereignty")

    c1, c2 = st.columns(2)
    c1.metric("External calls", sov["external_connections"])
    c2.metric("Blocked", sov["blocked_attempts"])
    backend = "docker" if sov["sandbox"]["docker"] else (
        "netns" if sov["sandbox"]["netns"] else "rlimit only")
    st.write(f"Sandbox: **{backend}**")
    st.write("Audit chain: " + ("**valid**" if sov["audit_chain_valid"]
                                else f"**BROKEN at {sov['audit_first_bad_index']}**"))

    if st.button("Attempt external call", use_container_width=True):
        r = api_post("/sovereignty/canary")
        (st.success if r["blocked"] else st.warning)(
            f"{r['verdict']} \u2014 {r['detail']}")
        if not r["blocked"]:
            st.caption("Expected on a machine with internet. Run "
                       "`sudo ./scripts/04_lockdown.sh on` to enforce denial.")

    st.divider()
    st.subheader("Index")
    st.write(f"Graph: **{health['graph_nodes']}** nodes / **{health['graph_edges']}** edges")
    st.write(f"Vectors: **{health['vector_chunks']}** chunks")
    st.caption(f"profile `{health['profile']}` \u00b7 ontology `{S.ontology}`")


tabs = st.tabs(["Ask", "Search", "Graph", "Documents", "Models", "Audit"])

# ---------------------------------------------------------------------- ASK -
with tabs[0]:
    st.subheader("Run an agentic goal")
    st.caption("Plan \u2192 execute tools \u2192 critique \u2192 synthesise. "
               "Typically 45\u2013120 s. Always open the generated file.")

    examples = [
        "Which units are affected if P-101A is isolated, and draft an approval "
        "note recommending seal replacement.",
        "Summarise the P-101A inspection findings.",
        "What is downstream of E-204?",
    ]
    pick = st.selectbox("Example goals", ["(write my own)"] + examples)
    goal = st.text_area("Goal", value="" if pick.startswith("(") else pick, height=90)

    c1, c2, _ = st.columns([1, 1, 3])
    iters = c1.number_input("Max iterations", 1, 3, 1)
    steps = c2.number_input("Max steps", 1, 8, 4)

    if st.button("Run agent", type="primary", disabled=not goal.strip()):
        t0 = time.time()
        with st.spinner("Planning and executing..."):
            try:
                res = api_post("/ask", goal=goal, max_iterations=int(iters),
                               max_steps=int(steps))
            except requests.exceptions.RequestException as e:
                st.error(f"{type(e).__name__}: {e}")
                st.stop()
        st.success(f"Completed in {time.time()-t0:.0f}s \u00b7 {res['iterations']} iteration(s)")

        st.markdown("### Answer")
        st.markdown(res["answer"] or "_no answer produced_")

        st.markdown("### Steps")
        for s in res["steps"]:
            icon = "\u2705" if s["ok"] else "\u274C"
            with st.expander(f"{icon} `{s['tool']}` \u00b7 {s['ms']} ms \u00b7 {s['intent']}",
                             expanded=not s["ok"]):
                st.caption("Arguments actually passed (after coercion)")
                st.json(s["arguments"], expanded=False)
                st.caption("Raw tool output")
                st.code(s["output"])

        if res["decisions"]:
            st.markdown("### Model routing")
            st.dataframe(pd.DataFrame(res["decisions"])[
                ["task", "chosen", "tag", "reason", "evicted", "load_ms"]],
                use_container_width=True, hide_index=True)

        st.markdown("### Artefacts")
        if not res["artifacts"]:
            st.info("No file was generated.")
        for a in res["artifacts"]:
            p = Path(a)
            if not p.exists():
                st.error(f"{a} \u2014 reported but missing")
                continue
            st.write(f"**{p.name}** \u00b7 {p.stat().st_size:,} bytes")
            st.download_button("Download", p.read_bytes(), p.name,
                               key=f"dl-{p.name}")
            try:
                if p.suffix == ".xlsx":
                    st.dataframe(pd.read_excel(p), use_container_width=True)
                elif p.suffix == ".docx":
                    from docx import Document
                    st.text("\n".join(x.text for x in Document(p).paragraphs
                                      if x.text.strip())[:4000])
            except Exception as e:
                st.warning(f"Could not preview: {e}")

# ------------------------------------------------------------------- SEARCH -
with tabs[1]:
    st.subheader("Retrieval only")
    st.caption("Instant. Shows which of the four modes fired and why.")
    q = st.text_input("Query", placeholder="e.g. which units are affected if P-101A is isolated")
    k = st.slider("Passages", 1, 20, 6)
    if st.button("Retrieve", disabled=not q.strip()):
        t0 = time.time()
        res = api_get("/search", q=q, k=k)
        st.success(f"mode **{res['mode']}** in {(time.time()-t0)*1000:.0f} ms")
        st.caption(f"Why: {res['why']}" +
                   (f" \u00b7 anchors: {res['anchors']}" if res.get("anchors") else ""))

        nodes = res.get("nodes") or []
        if nodes:
            st.markdown("#### Graph nodes reached")
            st.dataframe(pd.DataFrame(nodes), use_container_width=True, hide_index=True)

        for c in res.get("communities") or []:
            with st.expander(f"Theme: {c['title']} ({c['size']} members)"):
                st.write(c["summary"])

        st.markdown("#### Passages")
        for c in res.get("chunks") or []:
            with st.expander(f"[{c.get('doc_id','?')} p.{c.get('page','?')}]"):
                st.text(c["text"][:2500])

# -------------------------------------------------------------------- GRAPH -
with tabs[2]:
    st.subheader("Knowledge graph")
    st.caption("Traversal is what vector search cannot do. Compare hop depths.")
    c1, c2 = st.columns([3, 1])
    ent = c1.text_input("Entity", placeholder="e.g. P-101A, or a person's name")
    hops = c2.slider("Hops", 1, 4, 2)
    if st.button("Traverse", disabled=not ent.strip()):
        res = api_get("/graph/neighborhood", entity=ent.strip(), hops=hops)
        if not res["found"]:
            st.warning(f"No entity matching '{ent}'. Try the Search tab to find names.")
        else:
            anchor = res["anchor"]
            rows = res["rows"]
            st.write(f"Anchor: **{anchor['name']}** (`{anchor['type']}`)")
            st.write(f"**{len(rows)}** entities within {hops} hop(s)")
            if rows:
                st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
            if res["indirect_only"]:
                st.info(f"{len(res['indirect_only'])} of these are reachable only beyond "
                        "one hop \u2014 indirect dependencies vector search would miss: "
                        + ", ".join(res["indirect_only"][:8]))

    with st.expander("Active ontology"):
        st.write(f"**{len(NODE_TYPES)}** node types \u00b7 **{len(REL_PAIRS)}** relations")
        st.dataframe(pd.DataFrame(
            [{"node type": n, "attributes": ", ".join(a)} for n, a in NODE_TYPES.items()]),
            use_container_width=True, hide_index=True)
    with st.expander("Node counts"):
        st.dataframe(pd.DataFrame(sorted(health["graph"].items(), key=lambda x: -x[1]),
                                  columns=["type", "count"]),
                     use_container_width=True, hide_index=True)

# ---------------------------------------------------------------- DOCUMENTS -
with tabs[3]:
    st.subheader("Corpus")
    INDEXED = {".pdf", ".txt", ".md"}
    files = sorted(p for p in S.corpus_dir.rglob("*") if p.is_file())
    if files:
        st.dataframe(pd.DataFrame([{
            "file": str(p.relative_to(S.corpus_dir)),
            "KB": round(p.stat().st_size / 1024, 1),
            "indexed": "yes" if p.suffix.lower() in INDEXED else "NO - skipped",
        } for p in files]), use_container_width=True, hide_index=True)
        if any(p.suffix.lower() not in INDEXED for p in files):
            st.warning("Files marked *NO - skipped* are not picked up by `make index`. "
                       "Images are readable on demand via the vision tools, or convert "
                       "them to PDF to have them indexed.")
    else:
        st.info(f"No documents in {S.corpus_dir}")

    up = st.file_uploader("Add documents", type=["pdf", "txt", "md"],
                          accept_multiple_files=True)
    if up:
        for f in up:
            (S.corpus_dir / f.name).write_bytes(f.getbuffer())
        st.success(f"Copied {len(up)} file(s). Stop the API, run `make index`, restart.")

    st.divider()
    st.subheader("Generated artefacts")
    arts = sorted(p for p in S.artifacts_dir.glob("*") if p.is_file())
    if not arts:
        st.info("Nothing generated yet.")
    for p in arts:
        with st.expander(f"{p.name} \u00b7 {p.stat().st_size:,} bytes"):
            st.download_button("Download", p.read_bytes(), p.name, key=f"a-{p.name}")
            try:
                if p.suffix == ".xlsx":
                    st.dataframe(pd.read_excel(p), use_container_width=True)
                elif p.suffix == ".docx":
                    from docx import Document
                    st.text("\n".join(x.text for x in Document(p).paragraphs
                                      if x.text.strip())[:4000])
            except Exception as e:
                st.warning(f"Preview failed: {e}")

# ------------------------------------------------------------------- MODELS -
with tabs[4]:
    st.subheader("Model catalogue")
    st.caption("Adding a model is a block in config/models.yaml \u2014 no code change.")
    m = api_get("/models")
    st.dataframe(pd.DataFrame([{
        "id": s["id"], "tag": s["tag"], "VRAM GB": s["vram_gb"], "ctx": s["ctx"],
        "priority": s["priority"], "capabilities": ", ".join(s["capabilities"]),
        "resident": "yes" if s["id"] in m["resident"] else "",
    } for s in m["catalogue"]]), use_container_width=True, hide_index=True)

    st.write(f"Budget **{m['budget_gb']} GB** \u00b7 resident **{m['used_gb']:.1f} GB**")
    if m["decisions"]:
        st.markdown("#### Routing decisions (API session)")
        st.dataframe(pd.DataFrame(m["decisions"])[
            ["task", "capability", "tag", "reason", "evicted", "load_ms"]],
            use_container_width=True, hide_index=True)
    else:
        st.info("No model has been invoked yet in the API session.")

    st.divider()
    st.subheader("Sandbox check")
    st.caption("Rubric point 3 and half of point 5, verifiable here. Runs locally in "
               "this console's own process, independent of the API.")
    code = st.text_area("Python to execute in the sandbox", height=120, value=(
        "import socket\n"
        "socket.setdefaulttimeout(3)\n"
        "try:\n"
        "    socket.create_connection(('1.1.1.1', 53))\n"
        "    print('REACHED INTERNET - sovereignty broken')\n"
        "except Exception as e:\n"
        "    print('blocked:', type(e).__name__, e)\n"))
    if st.button("Execute in sandbox"):
        r = run_python(code)
        st.write(f"backend **{r.backend}** \u00b7 exit {r.exit_code}")
        if r.stdout:
            st.code(r.stdout)
        if r.stderr:
            st.code(r.stderr)
        if "REACHED INTERNET" in r.stdout:
            st.error("Sandbox leaked network access.")
        elif r.ok:
            st.success("Executed with no network reachability.")

# -------------------------------------------------------------------- AUDIT -
with tabs[5]:
    st.subheader("Tamper-evident audit log")
    n = st.slider("Records", 10, 300, 40)
    a = api_get("/audit", n=n)
    (st.success if a["chain_valid"] else st.error)(
        "Hash chain verified" if a["chain_valid"] else f"Chain broken at record {a['first_bad_index']}")
    recs = a["records"]
    if recs:
        st.dataframe(pd.DataFrame([{
            "time": time.strftime("%H:%M:%S", time.localtime(r["ts"])),
            "event": r["event"],
            "detail": json.dumps(r["payload"], default=str)[:160],
        } for r in reversed(recs)]), use_container_width=True, hide_index=True)
    else:
        st.info("Audit log is empty.")
